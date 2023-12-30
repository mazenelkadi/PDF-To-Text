import tkinter as tk  # Tkinter for creating GUI elements
from tkinter import filedialog  # Dialogs for file selection
import pdfminer.high_level  # Library for extracting text from PDFs
from docx import Document  # Library to work with Word documents
from openpyxl import Workbook  # Library to work with Excel files
import os  # Provides functions for interacting with the operating system


def with_pdfminer(pdf):
    """Extract text from a PDF file using pdfminer."""
    with open(pdf, 'rb') as file_handle:
        return pdfminer.high_level.extract_text(file_handle)


def save_text_to_file(text, output_file):
    """Save the extracted text to a file in TXT, DOCX, or XLSX format."""
    text = clean_text(text)  # Clean the text
    ext = os.path.splitext(output_file)[1]

    if ext == ".txt":
        with open(output_file, 'w', encoding='utf-8') as file:
            file.write(text)
    elif ext == ".docx":
        doc = Document()
        doc.add_paragraph(text)
        doc.save(output_file)
    elif ext == ".xlsx":
        wb = Workbook()
        ws = wb.active
        for line in text.splitlines():
            ws.append([line])
        wb.save(output_file)
    else:
        print("Unsupported file format")


def select_pdf():
    """Open a dialog to select a PDF file."""
    root = tk.Tk()
    root.withdraw()
    return filedialog.askopenfilename(title="Select PDF file", filetypes=[("PDF Files", "*.pdf")])


def clean_text(text):
    """Remove non-printable characters from the text."""
    return ''.join(char for char in text if char.isprintable())


def save_file_dialog():
    """Open a dialog to save the extracted text."""
    root = tk.Tk()
    root.withdraw()
    return filedialog.asksaveasfilename(
        title="Save Converted Text",
        filetypes=[("Text Files", "*.txt"), ("Word Documents",
                                             "*.docx"), ("Excel Worksheets", "*.xlsx")],
        defaultextension=[("Text Files", "*.txt"), ("Word Documents",
                                                    "*.docx"), ("Excel Worksheets", "*.xlsx")]
    )


# Main execution
if __name__ == '__main__':
    pdf_file = select_pdf()  # Select a PDF file
    if pdf_file:
        output_file = save_file_dialog()  # Choose where to save the converted text
        if output_file:
            text = with_pdfminer(pdf_file)  # Extract text from the PDF
            save_text_to_file(text, output_file)  # Save the text
            print(f"PDF converted and saved as {output_file}!")
        else:
            print("File save cancelled.")
    else:
        print("No PDF file selected.")